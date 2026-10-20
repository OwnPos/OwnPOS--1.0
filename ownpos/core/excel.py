"""Excel export (formatted workbooks) and smart import (any reasonable Excel/CSV layout, checked preview).

Import reads .xlsx, .xls (old Excel, e.g. QuickBooks POS 2013 exports) and .csv. It finds the header row by
itself (title lines above it are fine), recognises columns by their names (QuickBooks POS names, distributor
price-list names, and this program's own templates), and the user confirms or corrects the guess before
anything is saved. UPCs that Excel turned into numbers (1.23457E+11) are restored exactly when possible and
flagged when they are not."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .lists import Entity, Field, ListService, ValidationError
from .barcode import normalize_code
from .money import to_dollars

HEADER_FILL = PatternFill("solid", fgColor="1F2A37")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
TOTAL_FONT = Font(bold=True)
THIN = Side(style="thin", color="D0D5DD")
MONEY_FMT = '#,##0.00;[Red]-#,##0.00'
QTY_FMT = '#,##0.###'
PCT_FMT = '0.0"%"'


def _cell_value(value, typ):
    if value is None:
        return None
    if typ == "money":
        return float(to_dollars(value))
    if typ in ("qty", "pct", "number", "hours"):
        try:
            return float(value)
        except (TypeError, ValueError):
            return value
    if typ == "int":
        try:
            return int(value)
        except (TypeError, ValueError):
            return value
    if typ in ("date", "datetime") and isinstance(value, str) and len(value) >= 10:
        # stored as ISO text; written as a real Excel date shown the U.S. way (formats set by the caller)
        try:
            return datetime.fromisoformat(value[:19]) if typ == "datetime" else datetime.fromisoformat(value[:10])
        except ValueError:
            return value
    return value


def _write_sheet(ws, title: str, columns: list[tuple], rows: list[dict], totals: bool = True,
                 company: str = "") -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    sub = f"{company}  ·  Printed {datetime.now():%m/%d/%Y %I:%M %p}"
    ws["A2"] = sub.strip(" ·")
    ws["A2"].font = Font(italic=True, color="667085")
    header_row = 4
    for ci, (key, label, typ) in enumerate(columns, start=1):
        c = ws.cell(row=header_row, column=ci, value=label)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=THIN)
    ws.row_dimensions[header_row].height = 30
    r = header_row
    for row in rows:
        r += 1
        style = row.get("style")
        for ci, (key, label, typ) in enumerate(columns, start=1):
            c = ws.cell(row=r, column=ci, value=_cell_value(row.get(key), typ))
            if typ == "money":
                c.number_format = MONEY_FMT
            elif typ == "qty":
                c.number_format = QTY_FMT
            elif typ == "pct":
                c.number_format = PCT_FMT
            elif typ in ("number", "hours"):
                c.number_format = '#,##0.00'
            elif typ == "date":
                c.number_format = "mm/dd/yyyy"
            elif typ == "datetime":
                c.number_format = "mm/dd/yyyy hh:mm AM/PM"
            if style in ("total", "grand", "header"):
                c.font = TOTAL_FONT
            if style == "grand":
                c.border = Border(top=THIN, bottom=Side(style="double"))
    last_data = r
    if totals and rows:
        r += 1
        ws.cell(row=r, column=1, value="Total").font = TOTAL_FONT
        for ci, (key, label, typ) in enumerate(columns, start=1):
            from .reports import NO_TOTAL_KEYS
            if typ in ("money", "qty", "int", "hours") and ci > 1 and key not in NO_TOTAL_KEYS:
                col = get_column_letter(ci)
                if any(rw.get("style") for rw in rows):
                    # the sheet has its own subtotal rows: add up only the detail rows (a formula would count
                    # the subtotals twice)
                    vals = [rw.get(key) for rw in rows if not rw.get("style")]
                    value = sum(v for v in vals if isinstance(v, (int, float)))
                    value = float(to_dollars(round(value))) if typ == "money" else value
                else:
                    value = f"=SUBTOTAL(9,{col}{header_row + 1}:{col}{last_data})"   # follows Excel's filters
                c = ws.cell(row=r, column=ci, value=value)
                c.font = TOTAL_FONT
                c.number_format = MONEY_FMT if typ == "money" else QTY_FMT
                c.border = Border(top=THIN)
    for ci, (key, label, typ) in enumerate(columns, start=1):
        longest = max([len(str(label))] + [len(str(_cell_value(rw.get(key), typ) or "")) for rw in rows[:500]])
        ws.column_dimensions[get_column_letter(ci)].width = min(max(10, longest + 2), 60)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{last_data}"
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape" if len(columns) > 6 else "portrait"
    ws.page_setup.fitToWidth = 1


def _sheet_name(title: str, used: set) -> str:
    base = "".join(ch for ch in title if ch not in '[]:*?/\\')[:28] or "Report"
    name, i = base, 2
    while name in used:
        name = f"{base[:25]} {i}"
        i += 1
    used.add(name)
    return name


def export_table(path: str | Path, title: str, columns: list[tuple], rows: list[dict], totals: bool = True,
                 company: str = "", params: dict | None = None, notes: str = "") -> Path:
    """columns: [(key, label, type)]. Money values are cents and are written as dollars."""
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_name(title, set())
    _write_sheet(ws, title, columns, rows, totals, company)
    info = wb.create_sheet("Report info")
    info["A1"], info["A1"].font = "Report", TOTAL_FONT
    info["B1"] = title
    info["A2"], info["B2"] = "Created", f"{datetime.now():%Y-%m-%d %H:%M:%S}"
    info["A3"], info["B3"] = "Rows", len(rows)
    i = 4
    for k, v in (params or {}).items():
        info.cell(row=i, column=1, value=k)
        info.cell(row=i, column=2, value=str(v))
        i += 1
    if notes:
        info.cell(row=i + 1, column=1, value="Notes").font = TOTAL_FONT
        info.cell(row=i + 1, column=2, value=notes)
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 60
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_book(path: str | Path, results: list, company: str = "") -> Path:
    """Several reports in one workbook, one sheet each."""
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    for res in results:
        ws = wb.create_sheet(_sheet_name(res.title, used))
        _write_sheet(ws, res.title, res.columns, res.rows, res.totals, company)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_report(path, result, company: str = "") -> Path:
    return export_table(path, result.title, result.columns, result.rows, result.totals, company,
                        result.params, result.notes)


# ---------------------------------------------------------------- lists
def _export_value(svc: ListService, f: Field, value):
    if f.type == "money":
        return float(to_dollars(value or 0))
    if f.type == "bool":
        return "Yes" if value else "No"
    if f.type == "ref" and value not in (None, ""):
        for v, lbl in svc.choices(f.ref):
            if str(v) == str(value):
                return lbl if f.ref not in ("unit", "tax_code") else v
    if f.type == "choice" and value is not None:
        return value
    return value


def export_list(path, svc: ListService, ent: Entity, include_inactive: bool = False) -> Path:
    cols = [(f.key, f.label, "text") for f in ent.fields]
    rows = []
    for r in svc.rows(ent, include_inactive):
        rows.append({f.key: _export_value(svc, f, r.get(f.key)) for f in ent.fields})
    typed_cols = [(f.key, f.label, "number" if f.type in ("money", "qty", "pct") else ("int" if f.type == "int" else "text"))
                  for f in ent.fields]
    for r in rows:              # codes are written as text so Excel never turns a UPC into 1.23E+11
        for f in ent.fields:
            if f.type == "code" and r.get(f.key) is not None:
                r[f.key] = str(r[f.key])
    return export_table(path, ent.title, typed_cols, rows, totals=False)


def write_template(path, svc: ListService, ent: Entity, with_rows: bool = False) -> Path:
    """An import template: one header row with the field labels, notes and dropdowns."""
    wb = Workbook()
    ws = wb.active
    ws.title = ent.title[:28]
    editable = [f for f in ent.fields if f.editable]
    for ci, f in enumerate(editable, start=1):
        c = ws.cell(row=1, column=ci, value=f.label + (" *" if f.required else ""))
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        ws.column_dimensions[get_column_letter(ci)].width = max(12, len(f.label) + 4)
        options = None
        if f.type == "choice":
            options = [str(v) for v, _ in f.choices]
        elif f.type == "bool":
            options = ["Yes", "No"]
        elif f.type == "ref":
            options = [str(lbl if f.ref not in ("unit", "tax_code", "price_level") else v) for v, lbl in svc.choices(f.ref)]
        if options and sum(len(o) + 1 for o in options) < 250:
            dv = DataValidation(type="list", formula1='"' + ",".join(o.replace('"', "") for o in options) + '"',
                                allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(ci)}2:{get_column_letter(ci)}2000")
    if with_rows:
        for ri, r in enumerate(svc.rows(ent), start=2):
            for ci, f in enumerate(editable, start=1):
                ws.cell(row=ri, column=ci, value=_export_value(svc, f, r.get(f.key)))
    ws.freeze_panes = "A2"
    help_ws = wb.create_sheet("How to use")
    help_ws["A1"] = f"Import template for {ent.title}"
    help_ws["A1"].font = TITLE_FONT
    lines = ["Fill one row per record. Columns marked * are required.",
             "Money is in dollars (e.g. 24.99). Yes/No columns accept Yes or No.",
             "Type UPCs as text (start the cell with ' ) so Excel keeps leading zeros.",
             f"Rows that match an existing record on {', '.join(ent.match_on) or 'the key'} update it; others are added.",
             "The app shows a preview with any errors before anything is saved."]
    for i, line in enumerate(lines, start=3):
        help_ws.cell(row=i, column=1, value=line)
    help_ws.column_dimensions["A"].width = 100
    path = Path(path)
    wb.save(path)
    return path


# Column names this import recognises by itself: QuickBooks Point of Sale item/customer/vendor exports,
# common distributor price-list words, and this program's own templates. Anything else is matched by hand
# in the column-matching screen, which can remember the layout for next time.
ALIASES = {
    "items": {
        "item_no": ["item number", "item #", "item no", "item no.", "item code", "sku", "product code", "code",
                    "plu", "item id"],
        "upc": ["upc", "upc code", "upc/ean", "ean", "barcode", "bar code", "gtin", "scan code"],
        "alt_lookup": ["alternate lookup", "alu", "alt lookup", "alternate code"],
        "name": ["item name", "item description", "description 1", "description", "desc", "product name", "product",
                 "brand", "name", "item"],
        "size": ["size", "bottle size", "volume", "item size", "pack size"],
        "attribute": ["attribute", "flavor", "flavour", "vintage", "style"],
        "department_id": ["department name", "department", "dept name", "dept", "category", "class", "group"],
        "unit": ["unit of measure", "uom", "unit", "units", "selling unit"],
        "price1": ["regular price", "retail price", "price 1", "price", "sell price", "selling price", "retail",
                   "sale price", "srp"],
        "price2": ["price 2", "price level 2"],
        "price3": ["price 3", "price level 3"],
        "price4": ["price 4", "price level 4"],
        "price5": ["price 5", "price level 5", "employee price", "staff price"],
        "msrp": ["msrp", "suggested retail", "list price"],
        "tax_code": ["tax code", "tax", "taxable"],
        "reorder_point": ["reorder point", "reorder point 1", "reorder pt", "reorder level", "min qty", "minimum"],
        "reorder_qty": ["order up to", "max qty", "maximum", "reorder qty", "restock level"],
        "vendor_id": ["vendor name", "vendor", "supplier", "distributor", "primary vendor"],
        "vendor_item_no": ["vendor item #", "vendor item number", "vendor code", "supplier item #", "dist item #"],
        "notes": ["item notes", "notes", "memo", "remarks"],
        "_open_qty": ["qty 1", "on hand", "on-hand qty", "qty on hand", "quantity on hand", "on hand qty", "qty",
                      "quantity", "stock", "in stock", "onhand"],
        "_open_cost": ["average unit cost", "avg cost", "average cost", "order cost", "unit cost", "cost",
                       "last cost"],
    },
    "customers": {
        "code": ["customer id", "customer #", "customer number", "account #", "account number", "code"],
        "name": ["customer name", "full name", "name", "customer", "company"],
        "company": ["company", "business name"],
        "phone": ["phone 1", "phone", "mobile", "cell", "phone number", "telephone"],
        "email": ["email", "e-mail", "email address"],
        "address": ["street", "street 1", "address", "address 1"],
        "city": ["city"],
        "state": ["state"],
        "zip": ["zip", "zip code", "postal code"],
        "price_level": ["price level", "price lvl"],
        "discount_pct": ["discount %", "discount percent", "discount"],
        "credit_limit": ["credit limit", "charge limit", "account limit"],
        "tax_exempt": ["tax exempt", "exempt"],
        "notes": ["notes", "memo"],
        "_opening": ["balance", "account balance", "opening balance", "balance due", "a/r balance"],
    },
    "vendors": {
        "code": ["vendor code", "vendor id", "vendor #", "code"],
        "name": ["vendor name", "company", "vendor", "name", "distributor"],
        "contact": ["contact", "contact name", "salesperson", "rep"],
        "phone": ["phone 1", "phone", "telephone", "phone number"],
        "email": ["email", "e-mail"],
        "address": ["street", "street 1", "address", "address 1"],
        "city": ["city"],
        "state": ["state"],
        "zip": ["zip", "zip code"],
        "account_no": ["account #", "account number", "our account #", "customer #"],
        "terms_days": ["terms", "terms days", "net days", "pay within"],
        "notes": ["notes", "memo"],
        "_opening": ["balance", "opening balance", "balance due", "a/p balance"],
    },
    "employees": {
        "username": ["login", "login name", "username", "user name", "user"],
        "full_name": ["full name", "employee name", "name"],
        "title": ["title", "job title", "position"],
        "role": ["role", "security group"],
        "phone": ["phone", "mobile", "cell"],
        "pay_rate": ["pay rate", "hourly rate", "rate", "wage"],
        "notes": ["notes"],
    },
}
AUTO_CREATE_REFS = {"department", "vendor", "size"}

# extra columns that are not list fields: applied only to rows that are NEW in this import
EXTRA_FIELDS = {
    "items": [Field("_open_qty", "Opening on-hand quantity", "qty"), Field("_open_cost", "Opening unit cost ($)", "money")],
    "customers": [Field("_opening", "Opening balance: customer owes us ($)", "money")],
    "vendors": [Field("_opening", "Opening balance: we owe them ($)", "money")],
}


def import_fields(ent: Entity) -> list[Field]:
    """Every column an import into this list can fill: the editable fields plus the opening extras."""
    return [f for f in ent.fields if f.editable] + EXTRA_FIELDS.get(ent.key, [])


@dataclass
class NewRef:
    """A department, vendor or size named in the file that does not exist yet."""
    ref: str
    name: str


@dataclass
class ImportRow:
    row_no: int
    values: dict
    action: str = "add"         # add / update / skip
    pk: object = None
    errors: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    extra: dict = field(default_factory=dict)     # opening quantity / cost / balance (new rows only)


def _load_rows(path) -> list[list]:
    path = str(path)
    if path.lower().endswith(".xls"):
        import xlrd
        book = xlrd.open_workbook(path)
        sh = book.sheet_by_index(0)
        out = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                else:
                    row.append(cell.value)
            out.append(row)
        return out
    if path.lower().endswith((".csv", ".txt")):
        import csv
        for enc in ("utf-8-sig", "cp1252"):
            try:
                with open(path, newline="", encoding=enc) as fh:
                    return [list(r) for r in csv.reader(fh)]
            except UnicodeDecodeError:
                continue
        raise ValueError("Could not read the CSV file's text encoding")
    wb = load_workbook(path, data_only=True, read_only=True)
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def _cell(raw, i):
    """Cell i of a row, or None when the row is shorter (missing trailing cells) or i is None."""
    if raw is None or i is None or i < 0 or i >= len(raw):
        return None
    v = raw[i]
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _blank(raw) -> bool:
    return raw is None or all(_cell(raw, i) is None for i in range(len(raw)))


def _norm(h) -> str:
    return " ".join(str(h or "").replace("*", "").replace("_", " ").replace(":", "").strip().lower().split())


def guess_mapping(ent: Entity, header: list[str], saved: dict | None = None) -> dict:
    """{field key: column index}. A remembered layout (field key -> heading text) wins; then the field's own
    label or key; then the common names in ALIASES. Each column is used once."""
    aliases = ALIASES.get(ent.key, {})
    m, used = {}, set()
    for key, heading in (saved or {}).items():
        h = _norm(heading)
        if h in header and header.index(h) not in used:
            m[key] = header.index(h)
            used.add(m[key])
    for f in import_fields(ent):
        if f.key in m:
            continue
        for name in [_norm(f.label), f.key.lower()] + [_norm(a) for a in aliases.get(f.key, [])]:
            if name in header and header.index(name) not in used:
                m[f.key] = header.index(name)
                used.add(m[f.key])
                break
    return m


def read_header(path, ent: Entity, saved: dict | None = None) -> tuple[int, list, list, dict]:
    """(header row number, the headings as written, up to 3 sample rows, guessed mapping). Files with title
    lines above the headings (like QuickBooks exports) work: the row matching most fields is the header."""
    rows = _load_rows(path)
    best = (0, {})
    for hi in range(min(30, len(rows))):
        m = guess_mapping(ent, [_norm(h) for h in (rows[hi] or [])], saved)
        if len(m) > len(best[1]):
            best = (hi, m)
    hi, m = best
    headings = [str(h).strip() if h is not None else "" for h in (rows[hi] if rows else [])]
    samples = [r for r in rows[hi + 1:hi + 12] if not _blank(r)][:3]
    return hi, headings, samples, m


def read_import(path, svc: ListService, ent: Entity, mapping: dict | None = None,
                header_row: int | None = None) -> list[ImportRow]:
    """Read the file into checked rows. mapping (field key -> column index) comes from the column-matching
    screen; without it the columns are recognised automatically."""
    rows = _load_rows(path)
    if not rows:
        return []
    if mapping is None or header_row is None:
        hi, _headings, _samples, guessed = read_header(path, ent)
        by_label = guessed if mapping is None else mapping
        header_row = hi if header_row is None else header_row
    else:
        by_label = mapping
    hi = header_row
    header = [_norm(h) for h in (rows[hi] or [])]
    first_i = header.index("first name") if "first name" in header else None
    last_i = header.index("last name") if "last name" in header else None
    out = []
    for i, raw in enumerate(rows[hi + 1:], start=hi + 2):
        if _blank(raw):
            continue
        ir = ImportRow(i, {})
        for f in import_fields(ent):
            if f.key not in by_label:
                continue
            val = _cell(raw, by_label[f.key])
            try:
                if val in (None, "") and not f.required:
                    continue
                stored = svc.coerce(f, val)
                if f.key.startswith("_"):
                    ir.extra[f.key] = stored
                else:
                    ir.values[f.key] = stored
            except ValidationError as exc:
                if f.type == "ref" and f.ref in AUTO_CREATE_REFS and val not in (None, ""):
                    ir.values[f.key] = NewRef(f.ref, str(val))
                    ir.notes.append(f"new {f.ref} '{val}'")
                else:
                    ir.errors.append(str(exc))
        if ent.key == "customers" and not ir.values.get("name") and (first_i is not None or last_i is not None):
            parts = [str(_cell(raw, j)) for j in (first_i, last_i) if _cell(raw, j) is not None]
            if parts:
                ir.values["name"] = " ".join(parts)
        if ent.key == "customers" and "name" in ir.errors and ir.values.get("name"):
            ir.errors = [e for e in ir.errors if "Name" not in e]
        if ent.key == "items" and "upc" in by_label:
            raw_upc = _cell(raw, by_label["upc"])
            if isinstance(raw_upc, float) and raw_upc >= 1e15:
                ir.errors.append("UPC was stored by Excel as a rounded number; re-export it as text")
            elif isinstance(raw_upc, str) and ("e+" in raw_upc.lower()) and raw_upc.replace(".", "").replace("E+", "").replace("e+", "").isdigit():
                ir.errors.append(f"UPC '{raw_upc}' is in scientific notation, so digits are lost; re-export it as text")
            else:
                from .barcode import has_valid_check_digit
                code = normalize_code(raw_upc)
                if code.isdigit() and len(code) in (12, 13) and not has_valid_check_digit(code):
                    ir.notes.append(f"UPC {code}: the last digit is not a valid check digit; check it against the "
                                    "bottle")
        # match existing
        for key in ent.match_on:
            v = ir.values.get(key)
            if v in (None, "") or isinstance(v, NewRef):
                continue
            sql = f"SELECT {ent.pk} FROM {ent.table} WHERE lower(CAST({key} AS TEXT)) = lower(?)"
            if ent.where:
                sql += f" AND {ent.where}"
            hit = svc.db.one(sql, (v,))
            if hit is not None:
                ir.action, ir.pk = "update", hit[0]
                break
        if ir.action == "add":
            for f in ent.fields:
                if ent.key == "items" and f.key == "item_no" and f.key not in ir.values:
                    ir.notes.append("item number given automatically")
                    continue
                if f.required and f.key not in ir.values:
                    msg = f"{f.label} is required"
                    if msg not in ir.errors:
                        ir.errors.append(msg)
        elif ir.extra:
            ir.notes.append("already in the list: its opening amounts are NOT added again")
            ir.extra = {}
        ir.errors = [e for e in ir.errors if not (e.endswith("is required") and
                                                  any(f.label + " is required" == e and f.key in ir.values for f in ent.fields))]
        if ir.errors:
            ir.action = "skip"
        out.append(ir)
    return out


def _resolve_ref(svc: ListService, ref: NewRef, user_id: int, cache: dict):
    key = (ref.ref, ref.name.lower())
    if key in cache:
        return cache[key]
    from .lists import get_entity
    if ref.ref == "department":
        row = svc.db.one("SELECT id FROM departments WHERE lower(name) = lower(?)", (ref.name,))
        pk = row[0] if row else svc.save(get_entity("departments"), {"name": ref.name}, user_id)
    elif ref.ref == "size":
        code = ref.name.strip().upper().replace(" ", "")
        row = svc.db.one("SELECT code FROM lookups WHERE list_name = 'size' AND lower(code) = lower(?)", (code,))
        if row is None:
            svc.save(get_entity("lookup:size"), {"code": code, "name": ref.name.strip()}, user_id)
        pk = row[0] if row else code
    else:
        row = svc.db.one("SELECT id FROM vendors WHERE lower(name) = lower(?)", (ref.name,))
        pk = row[0] if row else svc.save(get_entity("vendors"), {"name": ref.name}, user_id)
    cache[key] = pk
    return pk


def commit_import(svc: ListService, ent: Entity, rows: list[ImportRow], user_id: int, services=None,
                  opening_date: str | None = None) -> dict:
    """Save the checked rows. With services given, the opening extras of NEW rows are posted too:
    items -> opening stock (quantity at the unit cost); customers/vendors -> opening balance on opening_date.
    Everything is one transaction: if the import stops half-way, nothing of it is saved."""
    added = updated = failed = 0
    errors = []
    cache: dict = {}
    openings = []
    with svc.db.tx():
        for ir in rows:
            if ir is None or ir.action == "skip" or not ir.values:
                continue
            # each row in its own savepoint: a row that fails half-way leaves nothing behind
            snapshot = dict(cache)
            svc.db.exec("SAVEPOINT import_row")
            try:
                values = {k: (_resolve_ref(svc, v, user_id, cache) if isinstance(v, NewRef) else v)
                          for k, v in ir.values.items()}
                if ent.key == "items" and ir.action == "add" and not values.get("item_no"):
                    values["item_no"] = svc.next_item_no()
                pk = svc.save(ent, values, user_id, ir.pk if ir.action == "update" else None)
                svc.db.exec("RELEASE SAVEPOINT import_row")
                added += ir.action == "add"
                updated += ir.action == "update"
                if ir.action == "add" and ir.extra:
                    openings.append((pk, ir.extra, ir.row_no))
            except Exception as exc:  # keep importing other rows
                svc.db.exec("ROLLBACK TO SAVEPOINT import_row")
                svc.db.exec("RELEASE SAVEPOINT import_row")
                cache.clear()
                cache.update(snapshot)
                failed += 1
                errors.append(f"Row {ir.row_no}: {exc}")
        opened = 0
        if services is not None and openings:
            if ent.key == "items":
                lines = []
                for pk, extra, _row in openings:
                    qty = float(extra.get("_open_qty") or 0)
                    if qty:
                        item = svc.db.one("SELECT kind, last_cost FROM items WHERE id = ?", (pk,))
                        if item["kind"] == "inventory":
                            lines.append((pk, qty, float(extra.get("_open_cost") or item["last_cost"] or 0)))
                if lines:
                    services.stock.opening(lines, user_id)
                    opened = len(lines)
            elif ent.key in ("customers", "vendors"):
                ptype = "customer" if ent.key == "customers" else "vendor"
                for pk, extra, _row in openings:
                    amt = int(extra.get("_opening") or 0)
                    if amt:
                        services.accounting.party_opening_balance(ptype, pk, amt, user_id, opening_date)
                        opened += 1
    return {"added": added, "updated": updated, "failed": failed, "errors": errors,
            "created_refs": len(cache), "openings": opened}


OPENING_ALIASES = {
    "sku": ["item #", "item no.", "item no", "item number", "sku", "code", "upc"],
    "qty": ["qty", "quantity", "on hand", "on-hand qty", "qty 1", "qty on hand", "counted"],
    "cost": ["unit cost", "cost", "average unit cost", "avg cost", "order cost"],
}


def read_opening_stock(path, db) -> tuple[list[tuple[int, float, int]], list[str]]:
    """Opening stock from Excel/CSV: item number (or UPC), quantity, unit cost (dollars). Returns (lines, problems).

    Blank rows are skipped; rows shorter than the header are read with the missing cells empty.
    An empty unit cost uses the item's current average (or last) cost; if the item has none the
    line is imported at cost 0 and a warning is added to problems."""
    from .money import to_cents as to_paisa
    rows = _load_rows(path)
    for hi in range(min(30, len(rows))):
        header = [_norm(h) for h in (rows[hi] or [])]
        idx = {k: next((header.index(a) for a in names if a in header), None) for k, names in OPENING_ALIASES.items()}
        if idx["sku"] is not None and idx["qty"] is not None:
            break
    else:
        return [], ["Could not find 'Item no.' and 'Qty' columns"]
    lines, problems = [], []
    for n, raw in enumerate(rows[hi + 1:], start=hi + 2):
        if _blank(raw):
            continue
        sku_v = _cell(raw, idx["sku"])
        if sku_v is None:
            problems.append(f"Row {n}: no item number")
            continue
        sku = normalize_code(sku_v)
        item = (db.one("SELECT id, cost_avg, last_cost FROM items WHERE lower(item_no) = lower(?)", (sku,))
                or db.one("SELECT id, cost_avg, last_cost FROM items WHERE upc = ?", (sku,)))
        if item is None:
            problems.append(f"Row {n}: item {sku} not found")
            continue
        qty_v = _cell(raw, idx["qty"])
        if qty_v is None:
            problems.append(f"Row {n}: item {sku} has no quantity")
            continue
        try:
            qty = float(str(qty_v).replace(",", ""))
        except ValueError:
            problems.append(f"Row {n}: quantity '{qty_v}' is not a number")
            continue
        if not qty:
            continue
        cost_v = _cell(raw, idx["cost"])
        if cost_v is not None:
            try:
                cost = to_paisa(cost_v)
            except (TypeError, ValueError, ArithmeticError):
                problems.append(f"Row {n}: unit cost '{cost_v}' is not an amount")
                continue
        else:
            cost = int(round(float(item["cost_avg"] or item["last_cost"] or 0)))
            if not cost:
                problems.append(f"Row {n}: item {sku} has no unit cost; imported at cost 0")
        lines.append((item["id"], qty, cost))
    return lines, problems
